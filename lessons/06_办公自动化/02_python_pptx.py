"""python-pptx 操作 PPT 演示文稿

演示 python-pptx：创建演示文稿、加幻灯片/文本框/表格、设字号、保存后遍历读取。

运行: uv run python lessons/06_办公自动化/02_python_pptx.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt


# 定位 sample 目录
SAMPLE_DIR = Path(__file__).parent / "sample"
SAMPLE_DIR.mkdir(exist_ok=True)
PPT_PATH = SAMPLE_DIR / "演示.pptx"


print("=" * 55)
print("1. 创建演示文稿：加幻灯片、文本框")
print("=" * 55)
prs = Presentation()
# 默认演示文稿宽 10 英寸、高 7.5 英寸（4:3），可用 prs.slide_width 改
print(f"  默认幻灯片尺寸: {prs.slide_width} x {prs.slide_height} EMU")
print(f"  换算成英寸: {prs.slide_width / 914400:.1f} x {prs.slide_height / 914400:.1f} 英寸")

# 加第一张幻灯片，用版式 0（标题幻灯片）
slide1 = prs.slides.add_slide(prs.slide_layouts[0])
# 版式 0 通常有两个占位符：标题和副标题
title_shape = slide1.placeholders[0]
subtitle_shape = slide1.placeholders[1]
title_shape.text = "自动化汇报"
subtitle_shape.text = "用 python-pptx 生成"
print("  第 1 张: 用 slide_layouts[0]（标题版式），填标题和副标题占位符")

# 加第二张幻灯片，用版式 1（标题和内容）或 6（空白）
slide2 = prs.slides.add_slide(prs.slide_layouts[6])  # 6 是空白版式
# 在空白幻灯片上手动加文本框
txBox = slide2.shapes.add_textbox(
    Inches(1), Inches(1), Inches(8), Inches(2)
)
tf = txBox.text_frame
tf.text = "这是手动加的文本框"
# 加第二段，设字号
p2 = tf.add_paragraph()
p2.text = "第二段，24 号字"
p2.font.size = Pt(24)
# 加第三段，设字号和加粗
p3 = tf.add_paragraph()
p3.text = "第三段，加粗 18 号"
p3.font.size = Pt(18)
p3.font.bold = True
print("  第 2 张: 用 slide_layouts[6]（空白版式），add_textbox 手动加文本框")
print("  -> text_frame.text 设第一段，add_paragraph() 加后续段落")

print("\n" + "=" * 55)
print("2. 加表格")
print("=" * 55)
# 在第二张幻灯片上加表格
rows, cols = 4, 3
table_shape = slide2.shapes.add_table(
    rows, cols, Inches(1), Inches(4), Inches(8), Inches(2)
)
table = table_shape.table
# 填表头
headers = ["项目", "耗时(天)", "状态"]
for j, h in enumerate(headers):
    table.cell(0, j).text = h
# 填数据
data = [
    ("文件批处理", "5", "已完成"),
    ("网络请求", "3", "已完成"),
    ("办公自动化", "7", "进行中"),
]
for i, row in enumerate(data, start=1):
    for j, val in enumerate(row):
        table.cell(i, j).text = val
print(f"  在第 2 张幻灯片加了 {rows}x{cols} 表格")
print("  -> add_table 返回的是 shape，要取 .table 才是表格对象")
print("  -> cell(i, j).text 填单元格，和 python-docx 类似")

print("\n" + "=" * 55)
print("3. 保存演示文稿")
print("=" * 55)
prs.save(PPT_PATH)
print(f"  已保存到: {PPT_PATH}")
print(f"  共 {len(prs.slides)} 张幻灯片")

print("\n" + "=" * 55)
print("4. 读取演示文稿：遍历幻灯片和形状")
print("=" * 55)
prs2 = Presentation(PPT_PATH)
print(f"  读取到 {len(prs2.slides)} 张幻灯片")
for s_idx, slide in enumerate(prs2.slides, start=1):
    print(f"\n  --- 幻灯片 {s_idx} ---")
    for shape in slide.shapes:
        # 形状有多种：文本框、表格、图片等
        if shape.has_text_frame:
            text = shape.text_frame.text
            print(f"    [文本框] {text}")
        elif shape.has_table:
            t = shape.table
            print(f"    [表格] {len(t.rows)} 行 x {len(t.columns)} 列:")
            for row in t.rows:
                cells = [cell.text for cell in row.cells]
                print(f"      {cells}")
        else:
            print(f"    [其他形状] 类型={shape.shape_type}")
print("  -> 遍历 slides，每个 slide 遍历 shapes")
print("  -> has_text_frame / has_table 判断形状类型，再取对应内容")

print("\n" + "=" * 55)
print("5. 坐标单位：Inches / Cm / Pt / EMU")
print("=" * 55)
from pptx.util import Cm, Emu
print(f"  Inches(1) = {Inches(1)} EMU")
print(f"  Cm(2.54)  = {Cm(2.54)} EMU  (2.54 厘米 = 1 英寸)")
print(f"  Pt(72)    = {Pt(72)} EMU   (72 磅 = 1 英寸)")
print("  -> PPT 内部用 EMU（English Metric Unit）存坐标，1 英寸 = 914400 EMU")
print("  -> 写代码用 Inches() / Cm() / Pt() 换算，不用自己算 EMU")

print("\n" + "=" * 55)
print("6. 常用版式说明")
print("=" * 55)
print("  slide_layouts[0]  标题幻灯片（标题 + 副标题）")
print("  slide_layouts[1]  标题和内容")
print("  slide_layouts[5]  仅标题")
print("  slide_layouts[6]  空白（最常用，自己往上加东西）")
print("  -> 不同模板的版式数量和含义可能不同，6 通常是空白")

print("\n" + "=" * 55)
print("完成!")
print("=" * 55)
print("要点:")
print("  1. Presentation() 新建，add_slide(slide_layouts[N]) 加幻灯片，N 是版式编号")
print("  2. 加内容用 shapes.add_textbox / add_table，坐标用 Inches() / Cm() 换算")
print("  3. 读取遍历 slides -> shapes，用 has_text_frame / has_table 判断类型")
print("  4. 坐标内部用 EMU，1 英寸 = 914400 EMU，写代码用 Inches/Cm/Pt 包装函数")