"""openpyxl 读写 —— 操作 .xlsx Excel 文件

演示 openpyxl：创建工作簿、逐格写、批量 append、遍历读、设置样式。
csv 模块搞不了 Excel 格式（.xlsx 是压缩包不是纯文本），必须用 openpyxl。

运行: uv run python lessons/02_表格数据/02_openpyxl读写.py
"""
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

# 定位 sample 目录
sample = Path(__file__).parent / "sample"
sample.mkdir(exist_ok=True)
xlsx_file = sample / "销售数据.xlsx"

# --- 准备测试数据 ---
sales = [
    ("笔记本", 12, 5999, "2026-09-01"),
    ("鼠标", 45, 89, "2026-09-01"),
    ("键盘", 30, 299, "2026-09-02"),
    ("显示器", 8, 1599, "2026-09-02"),
    ("U盘", 60, 79, "2026-09-03"),
    ("移动硬盘", 22, 459, "2026-09-03"),
    ("耳机", 38, 199, "2026-09-04"),
    ("摄像头", 15, 359, "2026-09-04"),
    ("路由器", 20, 249, "2026-09-05"),
    ("音箱", 10, 699, "2026-09-05"),
]

print("=" * 55)
print("1. 创建工作簿并写入数据")
print("=" * 55)
wb = Workbook()              # 新建工作簿（自带一个空工作表）
ws = wb.active               # 取活动工作表
ws.title = "销售明细"        # 改工作表名

# 写表头
headers = ["产品", "销量", "单价", "日期"]
ws.append(headers)           # append 追加一行（列表）
# 写数据行
for row in sales:
    ws.append(row)
print(f"  工作表名: {ws.title}")
print(f"  已写入 {len(sales)} 行数据 + 1 行表头")

print("\n" + "=" * 55)
print("2. 单格读写：ws['A1'].value")
print("=" * 55)
ws['E1'] = "销售额"          # 在 E1 写个新表头
print(f"  ws['A1'].value = {ws['A1'].value!r}")
print(f"  ws['B1'].value = {ws['B1'].value!r}")
print(f"  ws['E1'].value = {ws['E1'].value!r}  (刚写入的)")
# 单格写公式：E2 = B2 * C2（销量 * 单价）
for row_idx in range(2, len(sales) + 2):
    ws.cell(row=row_idx, column=5, value=f"=B{row_idx}*C{row_idx}")
print(f"  E2 写了公式: {ws['E2'].value}")

print("\n" + "=" * 55)
print("3. 遍历读取：ws.iter_rows()")
print("=" * 55)
print("  前 4 行数据:")
for row in ws.iter_rows(min_row=1, max_row=4, values_only=True):
    print(f"    {row}")
print("  -> values_only=True 直接取值，不返回 Cell 对象")

print("\n" + "=" * 55)
print("4. 设置样式：字体加粗 + 列宽")
print("=" * 55)
# 表头加粗
bold_font = Font(bold=True)
for col in range(1, 6):
    ws.cell(row=1, column=col).font = bold_font
# 设置列宽（按列字母）
ws.column_dimensions['A'].width = 12
ws.column_dimensions['B'].width = 8
ws.column_dimensions['C'].width = 8
ws.column_dimensions['D'].width = 12
ws.column_dimensions['E'].width = 10
print("  表头已加粗，列宽已设置")
print("  -> 样式保存在 .xlsx 里，用 Excel 打开能看到效果")

print("\n" + "=" * 55)
print("5. 保存到文件")
print("=" * 55)
wb.save(xlsx_file)
print(f"  已保存: {xlsx_file.name}")

print("\n" + "=" * 55)
print("6. 重新读取验证（load_workbook）")
print("=" * 55)
wb2 = load_workbook(xlsx_file)   # 重新打开
ws2 = wb2.active
print(f"  工作表名: {ws2.title}")
print(f"  总行数: {ws2.max_row}，总列数: {ws2.max_column}")
print("  全部数据:")
for row in ws2.iter_rows(values_only=True):
    print(f"    {row}")

print("\n" + "=" * 55)
print("完成!")
print("=" * 55)
print("要点:")
print("  1. .xlsx 必须用 openpyxl，csv 标准库搞不了 Excel 格式")
print("  2. ws.append(列表) 批量写，ws['A1'] 单格写，iter_rows 遍历读")
print("  3. openpyxl 能保留字体/列宽/公式，pandas 读 Excel 底层也调它")